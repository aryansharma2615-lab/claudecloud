import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
cfg = json.load(open("av_config_v8.json"))
sl = json.load(open("out/slicer_facts.json"))
wb = Workbook(); ws = wb.active; ws.title = "BOM"
F = lambda **k: Font(name="Arial", **k)
ws["A1"] = "Shawarma Servo Mount v1 — BOM (AV v8)"; ws["A1"].font = F(bold=True, size=13)
ws["A2"] = "Blue = input you can edit. Status OWNED = on the SP shelf: unit #1 costs $0, bulk columns stay priced for product maths."; ws["A2"].font = F(italic=True, size=9)
ws["A4"] = "PETG $/g"; ws["B4"] = 0.025; ws["B4"].font = F(color="0000FF"); ws["B4"].comment = Comment("Source: Digitmakers.ca PETG 1 kg listing, $25 / kg (checked 2026-09-13)", "AV v8")
ws["A5"] = "Machine $/h"; ws["B5"] = 0.35; ws["B5"].font = F(color="0000FF"); ws["B5"].comment = Comment("SP machine rate used by every SP AV (electricity + wear)", "AV v8")
hdr = ["Part", "Kind", "Qty", "Status", "Grams (SLICER)", "Print min (SLICER)", "Each (CAD $)", "Unit #1 out of pocket", "×10 each", "×50 each", "×100 each", "Line ×100", "Supplier", "Link", "Source / formula"]
r0 = 7
for i, h in enumerate(hdr, 1):
    c = ws.cell(r0, i, h); c.font = F(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="FF6B35"); c.alignment = Alignment(wrap_text=True)
r = r0 + 1
for p in cfg["parts"]:
    if p.get("bom") is False: continue
    kind = p["kind"]
    ws.cell(r, 1, p["label"]); ws.cell(r, 2, kind); ws.cell(r, 3, p.get("qty", 1)).font = F(color="0000FF")
    if kind == "printed":
        s = sl["parts"][p["id"]]
        ws.cell(r, 4, "PRINT"); ws.cell(r, 5, s["grams"]).font = F(color="0000FF"); ws.cell(r, 6, round(s["time_s"] / 60, 1)).font = F(color="0000FF")
        ws.cell(r, 7, f"=E{r}*$B$4+F{r}/60*$B$5")
        ws.cell(r, 8, f"=G{r}*C{r}")
        for col in (9, 10, 11): ws.cell(r, col, f"=G{r}")
        ws.cell(r, 13, "Digitmakers.ca PETG"); ws.cell(r, 14, "https://www.digitmakers.ca/collections/petg-1-75-mm")
        ws.cell(r, 15, f"{sl['slicer']} · Ender-3 S1 Pro profile · grams × $/g + minutes × $/h")
    else:
        c = p.get("cost") or {}
        ws.cell(r, 4, "OWNED" if c.get("have") else "BUY")
        ws.cell(r, 7, c.get("each", 0)).font = F(color="0000FF")
        ws.cell(r, 8, f'=IF(D{r}="OWNED",0,G{r}*C{r})')
        b = c.get("bulk") or {}
        for col, k in ((9, "10"), (10, "50"), (11, "100")): ws.cell(r, col, b.get(k, c.get("each", 0))).font = F(color="0000FF")
        ws.cell(r, 13, c.get("supplier", "")); ws.cell(r, 14, c.get("url") or ""); ws.cell(r, 15, c.get("formula", ""))
    ws.cell(r, 12, f"=K{r}*C{r}")
    for col in (7, 8, 9, 10, 11, 12): ws.cell(r, col).number_format = '$#,##0.00;($#,##0.00);"-"'
    r += 1
ws.cell(r, 1, "Total").font = F(bold=True)
for col, L in ((5, "E"), (6, "F"), (8, "H"), (12, "L")):
    ws.cell(r, col, f"=SUM({L}{r0+1}:{L}{r-1})").font = F(bold=True)
ws.cell(r, 13, "unit #1 = filament + time only (all hardware owned)")
for col in (8, 12): ws.cell(r, col).number_format = '$#,##0.00'
widths = [26, 9, 5, 9, 10, 10, 10, 12, 9, 9, 9, 10, 26, 44, 60]
for i, w in enumerate(widths, 1): ws.column_dimensions[chr(64 + i)].width = w
for row in ws.iter_rows(min_row=r0 + 1, max_row=r):
    for c in row:
        if c.font.color is None or c.font.color.rgb in (None, "FF000000"): c.font = F(bold=c.font.bold)
ws.freeze_panes = "B8"
wb.save("../../BOM_v8.xlsx"); print("rows", r - r0 - 1)
